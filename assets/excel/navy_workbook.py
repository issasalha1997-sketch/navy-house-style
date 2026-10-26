"""Navy house style for Excel workbooks built with openpyxl.

Use it as a module (copy it next to your builder and `import navy_workbook as nw`), or run it to build a
small demo workbook that shows the look:

    python navy_workbook.py demo.xlsx

Tested with openpyxl 3.1. Things openpyxl cannot do are noted where they come up: buttons (shapes with a
macro), solid data bars and dynamic-array formulas need Excel itself (see references/spreadsheets.md).
"""
import os
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Tahoma"
NAVY, INK, SOFT, WHITE = "243168", "1B1F2A", "2C3A78", "FFFFFF"
LINE, RULE = "D3D8EA", "C9CEDF"                     # row hairlines; tile and box borders
PICKER, TINT, BAR = "EEF0F7", "F4F5FA", "5B6AA6"
REF_TAB, FADED = "8F99C2", "7A82AE"
STATUS = {"done": "1E7B3C", "waiting": "8A5300", "with_others": "00618A", "stopped": "A4262C"}
LOOK_FILL, FILL_IN = "FFE699", "FFF4CC"
BANNER_FILL, BANNER_EDGE = "FFF4CE", "C27C0E"
LINK, LINK_FOLLOWED = "0563C1", "954F72"
DATE_FORMAT = "d mmm yyyy"
LINE_PT = 12.75                                     # Excel's line height for Tahoma 10 pt


def font(size=10, bold=False, color=INK, italic=False, **kw):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color, **kw)


def new_workbook():
    """A workbook whose own font is Tahoma 10, so empty cells, new rows and the width unit use it too.
    This sets two private openpyxl fields (_fonts, _named_styles): tested with openpyxl 3.1.x, so pin that version."""
    wb = Workbook()
    wb._fonts = IndexedList([Font(name=FONT, size=10, family=2)])
    wb._named_styles["Normal"].font = Font(name=FONT, size=10, family=2)
    return wb


def page_sheet(ws, widths, tab=NAVY):
    """A sheet that reads like a page: gridlines off, a narrow gutter column first. widths: {"A": 2, "B": 34, ...}"""
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_properties.tabColor = tab


def title(ws, cell, text, size=20, height=32):
    ws[cell] = text
    ws[cell].font = font(size, True, NAVY)
    ws[cell].alignment = Alignment(vertical="center")
    ws.row_dimensions[ws[cell].row].height = height


def subtitle(ws, cell, text):
    ws[cell] = text
    ws[cell].font = font(10, False, SOFT)


def note(ws, ref, text, lines=2):
    """A note under a box: 8.5 italic soft, merged, wrapped, top-aligned, fixed height (nobody types here)."""
    ws.merge_cells(ref)
    first = ws[ref.split(":")[0]]
    first.value = text
    first.font = font(8.5, False, SOFT, italic=True)
    first.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[first.row].height = lines * LINE_PT + 5


def rule_under(ws, row, first_col, last_col, style="medium", color=NAVY, height=None):
    for c in range(first_col, last_col + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style=style, color=color))
    if height:
        ws.row_dimensions[row].height = height


def section_title(ws, cell, text, size=13, height=24):
    ws[cell] = text
    ws[cell].font = font(size, True, NAVY)
    ws[cell].alignment = Alignment(vertical="bottom")
    ws.row_dimensions[ws[cell].row].height = height


def box_header(ws, row, first_col, labels, left_cols=(0,), height=28):
    """Dashboard table header: 9.5 bold navy, wrapped, thin navy underline; labels left, numbers right."""
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=first_col + j, value=lab)
        c.font = font(9.5, True, NAVY)
        c.alignment = Alignment(horizontal="left" if j in left_cols else "right", vertical="bottom", wrap_text=True)
        c.border = Border(bottom=Side(style="thin", color=NAVY))
    ws.row_dimensions[row].height = height


def box_row(ws, row, first_col, values, left_cols=(0,), size=11, bold_first=True, height=20):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=first_col + j, value=v)
        c.font = font(size, bold_first and j == 0, INK)
        c.alignment = Alignment(horizontal="left" if j in left_cols else "right", vertical="center")
    ws.row_dimensions[row].height = height


def total_row(ws, row, first_col, values, height=20):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=first_col + j, value=v)
        c.font = font(11, True, NAVY)
        c.alignment = Alignment(horizontal="left" if j == 0 else "right", vertical="center")
        c.border = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="thin", color=NAVY))
    ws.row_dimensions[row].height = height


def tile(ws, number_cell, caption_cell, value, caption, size=22):
    """A big number over a small caption, inside one thin box."""
    side = Side(style="thin", color=RULE)
    n, t = ws[number_cell], ws[caption_cell]
    n.value, t.value = value, caption
    n.font = font(size, True, NAVY)
    n.alignment = Alignment(horizontal="center", vertical="center")
    n.border = Border(left=side, right=side, top=side)
    t.font = font(8.5, False, SOFT)
    t.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    t.border = Border(left=side, right=side, bottom=side)


def picker(ws, ref, value=None):
    """The one input cell on a dashboard: pale fill, medium navy box, 13 bold navy, unlocked so it still works on the
    protected dashboard. ref: one cell, or a range on one row, which is merged first."""
    if ":" in ref:
        ws.merge_cells(ref)
        row = ws[ref][0]                                # read after merging: the merged cells take the box edges
    else:
        row = (ws[ref],)
    box = Side(style="medium", color=NAVY)
    for i, c in enumerate(row):
        c.fill = PatternFill("solid", fgColor=PICKER)
        c.border = Border(top=box, bottom=box, left=box if i == 0 else None, right=box if i == len(row) - 1 else None)
    first = row[0]
    first.value = value
    first.font = font(13, True, NAVY)
    first.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    first.protection = Protection(locked=False)


def banner(ws, ref, bold_part, rest, height=36):
    """A one-line notice at the top: amber fill, medium amber left edge, first sentence bold."""
    ws.merge_cells(ref)
    first = ws[ref.split(":")[0]]
    first.value = CellRichText(TextBlock(InlineFont(rFont=FONT, sz=9.5, b=True, color=INK), bold_part),
                               TextBlock(InlineFont(rFont=FONT, sz=9.5, color=INK), " " + rest))
    first.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)
    for row in ws[ref]:
        for c in row:
            c.fill = PatternFill("solid", fgColor=BANNER_FILL)
    first.border = Border(left=Side(style="medium", color=BANNER_EDGE))
    ws.row_dimensions[first.row].height = height


def group_labels(ws, row, groups, first_col=1, height=20):
    """Labels over groups of columns: 10 bold navy, centred across the group without merging, medium navy rule.
    groups: [("Who", 3), ("Tracking", 4)]"""
    col = first_col
    for label, span in groups:
        for k in range(span):
            c = ws.cell(row=row, column=col + k, value=label if k == 0 else None)
            c.font = font(10, True, NAVY)
            c.alignment = Alignment(horizontal="centerContinuous", vertical="center")
            c.border = Border(bottom=Side(style="medium", color=NAVY))
        col += span
    ws.row_dimensions[row].height = height


def header_row(ws, row, columns, first_col=1, height=26, dates=()):
    """The white-on-navy header row of a data table. columns: [("Name", 22), ("Email", 32), ...].
    Each column also gets Tahoma 10 (and the date format for dates) so new rows inherit the look."""
    for j, (label, width) in enumerate(columns):
        c = ws.cell(row=row, column=first_col + j, value=label)
        c.font = font(10, True, WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        cd = ws.column_dimensions[get_column_letter(first_col + j)]
        cd.width = width
        cd.font = font(10)
        if label in dates:
            cd.number_format = DATE_FORMAT
    ws.row_dimensions[row].height = height


def add_table(ws, name, ref):
    """An Excel table with filter buttons and no stripes; the navy header comes from header_row()."""
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=False)
    ws.add_table(t)
    return t


def hairlines(ws, ref, test):
    """A thin pale-navy line under every row that holds data. test: e.g. 'COUNTA($A4:$H4)>0'"""
    ws.conditional_formatting.add(ref, FormulaRule(formula=[test], border=Border(bottom=Side(style="thin", color=LINE))))


def status_colours(ws, ref, cell, mapping):
    """Status shown by font colour only. mapping: {"Active": ("done", True, False), ...}: (role, bold, italic)."""
    for value, (role, bold, italic) in mapping.items():
        ws.conditional_formatting.add(ref, FormulaRule(formula=['%s="%s"' % (cell, value)],
                                                       font=Font(color=STATUS[role], bold=bold, italic=italic)))


def faded_rows(ws, ref, test):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[test], font=Font(color=FADED)))


def needs_a_look(ws, ref, test):
    """Amber fill on the exact cell that needs fixing; use the same test as the dashboard count."""
    ws.conditional_formatting.add(ref, FormulaRule(formula=[test], fill=PatternFill(bgColor=LOOK_FILL, fill_type="solid")))


def count_red_above_zero(ws, ref, cell):
    ws.conditional_formatting.add(ref, FormulaRule(formula=["N(%s)>0" % cell], font=Font(color=STATUS["stopped"], bold=True)))


def list_rule(ws, ref, source, strict, message, title="Not on the list", prompt=None):
    """Drop-down: strict (Stop) only where a count depends on the value, a warning elsewhere."""
    d = DataValidation(type="list", formula1="=" + source, allow_blank=True, showErrorMessage=True,
                       errorStyle="stop" if strict else "warning", errorTitle=title, error=message,
                       showInputMessage=bool(prompt), prompt=prompt)
    d.add(ref)
    ws.add_data_validation(d)


def date_rule(ws, ref, up_to_today=False):
    """Dates from 2020; up to today for things that already happened."""
    upper = "TODAY()" if up_to_today else "DATE(2040,12,31)"
    d = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2=upper, allow_blank=True,
                       showErrorMessage=True, errorTitle="Type a date",
                       error="Type a date, for example 16 Mar 2026, or press Ctrl+; for today.")
    d.add(ref)
    ws.add_data_validation(d)


def worked_out(ws, ref, message="This cell is worked out from the list. Change the values there."):
    """Formula cells refuse typing (a custom FALSE rule) without protecting the sheet."""
    d = DataValidation(type="custom", formula1="FALSE", allow_blank=True, showErrorMessage=True,
                       errorTitle="Worked out automatically", error=message)
    d.add(ref)
    ws.add_data_validation(d)


def print_setup(ws, landscape=True, title_rows=None):
    """A4, fit to one page wide, quiet centred footer, header row repeated on each page."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top, ws.page_margins.bottom = 0.6, 0.7
    ws.oddFooter.center.text = "&A, page &P of &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.font = "Tahoma,Regular"
    if title_rows:
        ws.print_title_rows = title_rows


def finish(path):
    """After wb.save(path): theme fonts to Tahoma (so styles Excel creates later, such as links, use it) and the
    theme link colours to #0563C1 / #954F72 (otherwise links added later come out #0000FF)."""
    handle, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(handle)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/theme/theme1.xml":
                xml = data.decode("utf-8")
                for part in ("majorFont", "minorFont"):
                    xml, n = re.subn(r'(<a:%s>\s*<a:latin typeface=")[^"]*(")' % part, r"\g<1>%s\g<2>" % FONT, xml)
                    assert n == 1, part
                xml, n1 = re.subn(r'(<a:hlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>" + LINK, xml)
                xml, n2 = re.subn(r'(<a:folHlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>" + LINK_FOLLOWED, xml)
                assert n1 == 1 and n2 == 1, "theme link colours not found"
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def demo(path):
    """A small register: Dashboard (first), Tracker (plain data in a table), Lists (drop-down values)."""
    wb = new_workbook()
    statuses = [("Waiting", "waiting", False, False), ("On hold", "waiting", True, False), ("With supplier", "with_others", False, False),
                ("Active", "done", True, False), ("Stopped", "stopped", False, False)]

    # Lists: drop-down values, a reference tab
    ls = wb.create_sheet("Lists")
    ls.sheet_properties.tabColor = REF_TAB
    title(ls, "B1", "Lists", size=14, height=28)
    ls["B2"] = "Values for the drop-downs. New teams go at the end of the Team list."
    ls["B2"].font = font(9, False, SOFT, italic=True)
    ls["B3"] = "Fixed: the Dashboard counts each status."
    ls["B3"].font = font(8.5, True, STATUS["stopped"], italic=True)
    ls["B3"].alignment = Alignment(wrap_text=True, vertical="bottom")
    ls.row_dimensions[3].height = 30
    header_row(ls, 4, [("Status", 22)], first_col=2)
    header_row(ls, 4, [("Team", 22)], first_col=4)
    for i, (s, *_r) in enumerate(statuses, start=5):
        ls.cell(row=i, column=2, value=s).font = font(10)
    for i, t in enumerate(["Data", "Platform", "Sales"], start=5):
        ls.cell(row=i, column=4, value=t).font = font(10)
    ls.column_dimensions["A"].width = 2
    ls.column_dimensions["C"].width = 2

    # Tracker: plain typeable data only
    ws = wb["Sheet"]
    ws.title = "Tracker"
    ws.sheet_properties.tabColor = NAVY
    cols = [("Name", 22), ("Email", 32), ("Team", 24), ("Status", 18), ("Requested", 13), ("Done on", 13), ("Notes", 60)]
    group_labels(ws, 1, [("Who", 3), ("Tracking", 3), ("Notes", 1)])
    header_row(ws, 2, cols, dates=("Requested", "Done on"))
    rows = [("Ana Lima", "ana@example.org", "Platform", "Active", date(2026, 9, 2), date(2026, 9, 9), ""),
            ("Ben Okafor", "ben@example.org", "Data", "Waiting", date(2026, 10, 14), None, "Asked on 14 Oct"),
            ("Chloe Martin", "chloe.example.org", "Sales", "On hold", date(2026, 10, 1), None, "Email to check"),
            ("Dev Patel", "dev@example.org", "Platform", "Stopped", date(2026, 8, 20), date(2026, 8, 30), "Left the team")]
    for r, values in enumerate(rows, start=3):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font(10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 7))
            if isinstance(v, date):
                cell.number_format = DATE_FORMAT
    last = 2 + len(rows)
    add_table(ws, "tblTracker", "A2:G%d" % last)
    ws.freeze_panes = "C3"
    full = "A3:G2000"
    hairlines(ws, full, "COUNTA($A3:$D3)>0")
    status_colours(ws, "D3:D2000", "$D3", {s: (role, b, i) for s, role, b, i in statuses})
    faded_rows(ws, full, '$D3="Stopped"')
    needs_a_look(ws, "B3:B2000", 'AND($B3<>"",ISERROR(SEARCH("@",$B3)))')
    list_rule(ws, "D3:D2000", "Lists!$B$5:$B$9", True, "Pick a status from the list (the Dashboard counts use it).")
    list_rule(ws, "C3:C2000", "Lists!$D$5:$D$50", False, "This team is not on the Lists tab yet. Click Yes to keep it anyway.")
    date_rule(ws, "E3:F2000", up_to_today=True)

    # Dashboard: first sheet, reads like a page
    ds = wb.create_sheet("Dashboard", 0)
    page_sheet(ds, {"A": 2, "B": 30, "C": 14, "D": 14, "E": 14, "F": 14, "G": 2})
    banner(ds, "B1:F1", "Type in the Tracker; this page counts by itself.",
           "It is protected, so nothing here can be typed over by mistake.")
    ds.row_dimensions[2].height = 10
    title(ds, "B3", "Requests at a glance")
    subtitle(ds, "B4", "Counts work themselves out from the Tracker. Nothing here needs typing.")
    ds.row_dimensions[5].height = 46                     # the button row (buttons are added in Excel)
    rule_under(ds, 6, 2, 6, height=6)
    ds.row_dimensions[7].height = 12
    section_title(ds, "B8", "By status")
    box_header(ds, 9, 2, ["Status", "Lines", "Oldest request"], left_cols=(0,))
    for i, (s, *_r) in enumerate(statuses):
        r = 10 + i
        box_row(ds, r, 2, [s, None, None])
        ds.cell(row=r, column=3, value='=COUNTIFS(tblTracker[Status],$B%d)' % r)
        # functions newer than Excel 2007 need the _xlfn. prefix in the file, or Excel shows #NAME?
        ds.cell(row=r, column=4, value='=IF(C%d=0,"",_xlfn.MINIFS(tblTracker[Requested],tblTracker[Status],$B%d))' % (r, r))
        ds.cell(row=r, column=4).number_format = DATE_FORMAT
        ds.cell(row=r, column=4).font = font(11)
        ds.cell(row=r, column=3).font = font(11)
    total = 10 + len(statuses)
    total_row(ds, total, 2, ["All lines", "=SUM(C10:C%d)" % (total - 1), ""])
    hairlines(ds, "B10:D%d" % (total - 1), '$B10<>""')
    worked_out(ds, "C10:D%d" % total)
    ds.row_dimensions[total + 1].height = 16
    section_title(ds, "B%d" % (total + 2), "Needs a look")
    box_header(ds, total + 3, 2, ["What", "Lines"])
    box_row(ds, total + 4, 2, ["Email without an @", '=COUNTIFS(tblTracker[Email],"<>",tblTracker[Email],"<>*@*")'], size=10.5, bold_first=False, height=18)
    count_red_above_zero(ds, "C%d" % (total + 4), "C%d" % (total + 4))
    tile_row = total + 6
    tile(ds, "C%d" % tile_row, "C%d" % (tile_row + 1), '=COUNTIFS(tblTracker[Status],"Active")', "Active now")
    tile(ds, "D%d" % tile_row, "D%d" % (tile_row + 1), '=COUNTIFS(tblTracker[Status],"Waiting")+COUNTIFS(tblTracker[Status],"On hold")', "Waiting on us")
    tile(ds, "E%d" % tile_row, "E%d" % (tile_row + 1), '=COUNTIFS(tblTracker[Status],"With supplier")', "With the supplier")
    ds.row_dimensions[tile_row].height = 40
    ds.row_dimensions[tile_row + 1].height = 30
    ds.cell(row=tile_row, column=2, value="This month").font = font(13, True, NAVY)
    pick = tile_row + 3                                  # one team in detail: the picker and what answers to it
    ds.row_dimensions[pick].height = 26
    ds.cell(row=pick, column=2, value="Team").font = font(10, False, SOFT)
    ds.cell(row=pick, column=2).alignment = Alignment(horizontal="right", vertical="center", indent=1)
    picker(ds, "C%d:D%d" % (pick, pick), "Platform")
    list_rule(ds, "C%d" % pick, "Lists!$D$5:$D$50", True, "Pick a team from the list.")
    ds.cell(row=pick, column=5, value='=COUNTIFS(tblTracker[Team],$C$%d)' % pick).font = font(13, True, NAVY)
    ds.cell(row=pick, column=5).alignment = Alignment(horizontal="center", vertical="center")
    ds.cell(row=pick, column=6, value="lines for this team").font = font(9, False, SOFT, italic=True)
    ds.cell(row=pick, column=6).alignment = Alignment(vertical="center")
    worked_out(ds, "E%d" % pick)
    note(ds, "B%d:F%d" % (pick + 2, pick + 2),
         "Notes: a line counts once per status. Stopped lines stay in the Tracker with their date and reason.")
    print_setup(ds, landscape=False)
    print_setup(ws, title_rows="1:2")
    ds.protection.sheet = True                           # read-only counts; the Tracker stays open
    wb.save(path)
    finish(path)


if __name__ == "__main__":
    demo(sys.argv[1] if len(sys.argv) > 1 else "navy-demo.xlsx")
