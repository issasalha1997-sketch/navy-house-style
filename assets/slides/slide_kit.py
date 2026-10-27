"""Slide kit: native, editable PowerPoint slides in the house style, built with python-pptx (green by default;
Deck(theme="navy") for navy).

Everything it draws is a real PowerPoint object (shapes, lines, text boxes, icons PowerPoint can recolour), so the
slides stay editable. Text is measured with the real font before it is placed, so a line that would not fit stops
the build instead of spilling. Tested with python-pptx 1.0.x, Pillow 10+ and PyMuPDF 1.24+ (pip install python-pptx
pillow pymupdf). Run demo_deck.py to see every pattern.

Units are inches from the top-left corner of a 13.333 x 7.5 in (16:9) slide. Colours are hex strings without "#".
"""
import hashlib
import io
import math
import os
import pathlib
import re
import warnings
from datetime import datetime, timezone

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
ICONS = HERE / "icons"

# ---------------------------------------------------------------- canvas, grid, colours, type
SW, SH = 13.333333, 7.5          # 16:9 slide
L, R = 0.7, 12.63                # left and right margins: content runs from L to R

WHITE, INK = "FFFFFF", "000000"


class Palette:
    """The deck's colours, read at drawing time, so use_theme() changes everything drawn after it."""

    def __init__(self, **colours):
        self.__dict__.update(colours)


THEMES = {
    # Green (default): MongoDB's Forest Green as the primary and its Mist as the pale surface; the other steps are
    # tints of the primary, chosen for contrast (white text on primary 6.8:1, on "now" 4.6:1; hairline 1.8:1).
    "green": dict(primary="00684A", surface="E3FCF7", hairline="A6CAC0", decor="2E836B", now="2E836B",
                  ramp=["00684A", "4C9580", "8CBBAE", "C7DED7"]),
    # Navy: the house style's navy (white text on it 12.3:1).
    "navy": dict(primary="243168", surface="EEF0F7", hairline="BCC3DE", decor="3A4A86", now="3A4A86",
                 ramp=["243168", "5B6AA6", "8F99C2", "C9CEDF"]),
}
P = Palette(**THEMES["green"])


def use_theme(theme):
    """theme: "green", "navy", or a dict with primary, surface, hairline, decor, now and ramp (hex, no '#')."""
    P.__dict__.update(THEMES[theme] if isinstance(theme, str) else theme)


def _lum(hex6):
    c = [int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(a, b):
    la, lb = _lum(a), _lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def readable(colour, size, bold=False, on=WHITE):
    """The colour itself when text in it is readable at this size (4.5:1, or 3:1 for 18 pt, or 14 pt bold);
    otherwise the primary. Category colours that fail move onto marks (bars, outlines), never small text."""
    large = size >= 18 or (bold and size >= 14)
    return colour if contrast(colour, on) >= (3.0 if large else 4.5) else P.primary


FONT = "Tahoma"
LANG = "en-GB"
NB = "\u00a0"                    # non-breaking space: holds names and pairs together across line ends

SQUARE_DOT = MSO_LINE_DASH_STYLE.SQUARE_DOT
DASH = MSO_LINE_DASH_STYLE.DASH
ADEC_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"


def rgb(hex6):
    return RGBColor.from_string(hex6)


def nb(text):
    """Join the words of a name with non-breaking spaces, so it never splits across lines (screen text only;
    never inside wording quoted from a source)."""
    return NB.join(text.split())


# ---------------------------------------------------------------- measuring text with the real font
FONT_NAMES = {"Tahoma": (["tahoma.ttf"], ["tahomabd.ttf", "tahoma bold.ttf"]),
              "Verdana": (["verdana.ttf"], ["verdanab.ttf", "verdana bold.ttf"]),
              "DejaVu Sans": (["dejavusans.ttf"], ["dejavusans-bold.ttf"])}


def _find_font(names):
    """A font file by name (case-insensitive): first in the folder named by SLIDE_FONT_DIR, then in the usual
    Windows, macOS and Linux font folders, searched to any depth. Returns a path or None."""
    roots = [os.environ.get("SLIDE_FONT_DIR", ""), os.path.join(os.environ.get("WINDIR", ""), "Fonts"),
             os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts"),
             "/Library/Fonts", "/System/Library/Fonts", os.path.expanduser("~/Library/Fonts"),
             "/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
             os.path.expanduser("~/.local/share/fonts")]
    wanted = {n.lower() for n in names}
    for root in roots:
        if root and os.path.isdir(root):
            for dirpath, _, files in os.walk(root):
                for f in files:
                    if f.lower() in wanted:
                        return os.path.join(dirpath, f)
    return None


FONT_REGULAR = FONT_BOLD = None
LINE_H = 1.207                   # line height as a multiple of the font size (Tahoma's); set from the font file
_fonts = {}


def _measure_with(name, regular=None, bold=None):
    global FONT_REGULAR, FONT_BOLD, LINE_H
    reg_names, bold_names = FONT_NAMES.get(name, ([name.lower() + ".ttf"], [name.lower() + " bold.ttf"]))
    reg, bld = regular or _find_font(reg_names), bold or _find_font(bold_names)
    if not (reg and bld):
        return False
    FONT_REGULAR, FONT_BOLD = reg, bld
    _fonts.clear()
    asc, desc = ImageFont.truetype(reg, 2048).getmetrics()
    LINE_H = (asc + desc) / 2048
    return True


def use_font(name, regular=None, bold=None):
    """Write AND measure text in this font (for example "Verdana" for a deck going to Google Slides). Finds the files
    by name, or takes their paths; stops with a clear message if they cannot be found."""
    global FONT
    if not _measure_with(name, regular, bold):
        raise RuntimeError(f"No font files for {name}: install the font, set SLIDE_FONT_DIR to the folder that holds "
                           f"them, or pass regular= and bold= paths.")
    FONT = name


if not _measure_with("Tahoma"):                      # slides are still written in Tahoma; only the measuring changes
    for _stand_in in ("Verdana", "DejaVu Sans"):
        if _measure_with(_stand_in):
            warnings.warn(f"Tahoma not found: text is measured with {_stand_in}, so check the renders closely "
                          f"(or set SLIDE_FONT_DIR to a folder with tahoma.ttf and tahomabd.ttf).")
            break


def _font(size_pt, bold):
    if not FONT_REGULAR:
        raise RuntimeError("No font to measure text with: install Tahoma (or Verdana), or set SLIDE_FONT_DIR.")
    key = (round(size_pt * 20), bold)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, key[0])
    return _fonts[key]


def text_width_in(text, size_pt, bold=False):
    return _font(size_pt, bold).getlength(text) / 20 / 72.0


def lines_height_in(n_lines, size_pt, spacing=1.0):
    return n_lines * size_pt * LINE_H * spacing / 72.0


def layout_lines(runs, width_in, safety=1.04, min_last=0.5, balance=False):
    """Break runs [(text, size, bold, colour), ...] into lines that fit width_in.
    Fewest lines first, then the most even lines; the last line reaches at least min_last of the width (no
    stranded word). balance=True evens every line (good for two-line labels). Words split only at normal spaces,
    so non-breaking spaces hold. Returns a list of lines, each a list of runs."""
    limit = width_in * 72 * 20 / safety
    toks = []
    for text, size, bold, colour in runs:
        for word in text.split(" "):
            if word:
                toks.append((word, (size, bold, colour)))
    if not toks:
        return []
    widths = [_font(st[0], st[1]).getlength(w) for w, st in toks]
    spaces = [_font(st[0], st[1]).getlength(" ") for _, st in toks]

    def line_w(i, j):
        return sum(widths[i:j]) + sum(spaces[i + 1:j])

    n, inf = len(toks), float("inf")
    best = [(0.0, 0, None)] + [(inf, inf, None)] * n           # (cost, lines, start of the last line)
    for j in range(1, n + 1):
        for i in range(j - 1, -1, -1):
            w = line_w(i, j)
            if w > limit and j - i > 1:
                break
            c_prev, l_prev, _ = best[i]
            if c_prev == inf:
                continue
            slack = limit - w
            if j == n and not balance:
                cost = (max(0.0, min_last * limit - w) * 4) ** 2
            elif j == n:
                cost = 0.9 * slack ** 2
            else:
                cost = slack ** 2
            if (l_prev + 1, c_prev + cost) < (best[j][1], best[j][0]):
                best[j] = (c_prev + cost, l_prev + 1, i)
    cuts, j = [], n
    while j > 0:
        i = best[j][2]
        cuts.append((i, j))
        j = i
    lines = []
    for i, j in reversed(cuts):
        merged = []
        for word, st in toks[i:j]:
            if merged and merged[-1][1:] == st:
                merged[-1] = (merged[-1][0] + " " + word,) + st
            else:
                merged.append((word if not merged else " " + word,) + st)
        lines.append(merged)
    return lines


def N(text, size, colour=INK):
    """A regular run."""
    return (text, size, False, colour)


def B(text, size, colour=INK):
    """A bold run."""
    return (text, size, True, colour)


# ---------------------------------------------------------------- the deck
class Deck:
    """A 16:9 presentation with one layout, the brand theme and the kit's drawing helpers."""

    def __init__(self, theme="green", font=None, lang=None, theme_name="House slides", accents=None):
        global LANG
        if font and font != FONT:
            use_font(font)
        LANG = lang or LANG
        use_theme(theme)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)
        sld_sz = self.prs.part._element.find(qn("p:sldSz"))
        if sld_sz is not None and "type" in sld_sz.attrib:
            del sld_sz.attrib["type"]                            # the template calls itself 4:3
        self._keep_only_layout("Title Only")
        self._widen_master()
        self._set_theme(theme_name, accents or P.ramp + [P.now, P.hairline])
        self._svg_parts = {}
        self._svg_count = 0

    # ---- theme and layout
    def _keep_only_layout(self, keep):
        master = self.prs.slide_masters[0]
        lst = master._element.find(qn("p:sldLayoutIdLst"))
        for el in list(lst):
            rid = el.get(qn("r:id"))
            if master.part.related_part(rid).slide_layout.name != keep:
                master.part.drop_rel(rid)
                lst.remove(el)
        self.layout = next(lay for lay in master.slide_layouts if lay.name == keep)

    def _widen_master(self):
        """The template's master and layout were drawn for a 4:3 slide: stretch their x positions and widths to
        16:9, so a title, footer, date or slide number added later in PowerPoint lands in the right place."""
        k = 12192000 / 9144000
        master = self.prs.slide_masters[0]
        for owner in [master] + list(master.slide_layouts):
            for sp in owner.shapes:
                xfrm = sp._element.find(".//" + qn("a:xfrm"))
                if xfrm is None:
                    continue
                off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
                if off is not None:
                    off.set("x", str(int(int(off.get("x")) * k)))
                if ext is not None:
                    ext.set("cx", str(int(int(ext.get("cx")) * k)))

    def _set_theme(self, name, accents):
        """Theme fonts and colours, so anything added later (or pasted with Keep Source Formatting) starts in
        the brand font and colours instead of the template's."""
        part = self.prs.slide_masters[0].part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        root.set("name", name)
        scheme = root.find(f".//{{{a}}}clrScheme")
        scheme.set("name", name)
        slots = {"dk2": P.primary, "lt2": P.surface, "hlink": "0563C1", "folHlink": "954F72"}
        for i, hex6 in enumerate(accents[:6], start=1):
            slots[f"accent{i}"] = hex6
        for slot, hex6 in slots.items():
            el = scheme.find(f"{{{a}}}{slot}")
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{{{a}}}srgbClr").set("val", hex6)
        for which in ("majorFont", "minorFont"):
            root.find(f".//{{{a}}}{which}/{{{a}}}latin").set("typeface", FONT)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ---- slides
    def slide(self, backing=True):
        """A new slide with a white background. backing=True adds a locked white rectangle at the back, so the slide
        keeps its white ground when it is pasted into someone else's deck (the default paste drops the slide's own
        background and its 'hide background graphics' setting)."""
        s = self.prs.slides.add_slide(self.layout)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(WHITE)
        s._element.set("showMasterSp", "0")
        if backing:
            bg = self.box(s, 0, 0, SW, SH, fill=WHITE, name="White background (locked)")
            locks = etree.SubElement(bg._element.find(".//" + qn("p:cNvSpPr")), qn("a:spLocks"))
            for attr in ("noGrp", "noRot", "noMove", "noResize", "noEditPoints", "noAdjustHandles",
                         "noChangeArrowheads", "noChangeShapeType", "noTextEdit"):
                locks.set(attr, "1")                             # never noSelect: people must be able to delete it
            tree = bg._element.getparent()
            tree.remove(bg._element)
            tree.insert(2, bg._element)                          # right after nvGrpSpPr and grpSpPr: the very back
        return s

    # ---- text
    def _runs(self, p, runs):
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "0")
        ppr.set("indent", "0")
        for text, size, bold, colour in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(size)
            f.bold = bold
            f.italic = False
            f.color.rgb = rgb(colour)
            rpr = r._r.get_or_add_rPr()
            rpr.set("lang", LANG)
            rpr.set("cap", "none")                               # a host deck's ALL CAPS title style cannot leak in
            rpr.set("spc", "0")
            rpr.find(qn("a:latin")).addprevious(etree.Element(qn("a:effectLst")))

    def _break(self, p, size, bold):
        rpr = p._p.add_br().get_or_add_rPr()
        rpr.set("lang", LANG)
        rpr.set("sz", str(int(size * 100)))
        rpr.set("b", "1" if bold else "0")
        rpr.set("cap", "none")
        etree.SubElement(rpr, qn("a:latin")).set("typeface", FONT)

    @staticmethod
    def _frame(tf, anchor, margin=0.0):
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(margin))
        tf.vertical_anchor = anchor

    def _para_format(self, p, align, spacing):
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)

    def text_lines(self, s, x, y, w, h, lines, align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP, name=None):
        """One paragraph whose line ends are set by hand (explicit breaks), so PowerPoint never re-wraps it."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        self._frame(tb.text_frame, anchor)
        p = tb.text_frame.paragraphs[0]
        self._para_format(p, align, spacing)
        for i, line in enumerate(lines):
            if i:
                size = max(r[1] for r in line)
                self._break(p, size, line[0][2])
            self._runs(p, line)
        return tb

    def one(self, s, x, y, w, h, text, size, bold=False, colour=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            name=None):
        """A one-line label. Stops the build if it would not fit on one line."""
        assert text_width_in(text, size, bold) * 1.05 < w + 1e-6, f"{name or text!r} does not fit in {w:.2f} in"
        return self.text_lines(s, x, y, w, h, [[(text, size, bold, colour)]], align=align, anchor=anchor, name=name)

    def para(self, s, x, y, w, runs, max_lines, name, align=PP_ALIGN.LEFT, balance=False, spacing=1.0):
        """Wrap runs to width w (measured), refuse more than max_lines, size the box to the text. Returns
        (shape, height)."""
        lines = layout_lines(runs, w, balance=balance)
        assert len(lines) <= max_lines, f"{name}: needs {len(lines)} lines, room for {max_lines}"
        size = max(r[1] for r in runs)
        h = lines_height_in(len(lines), size, spacing) + 0.04
        return self.text_lines(s, x, y, w, h, lines, align=align, spacing=spacing, name=name), h

    def sentence_lines(self, s, x, y, w, sentences, size, colour=INK, name="Lead"):
        """One sentence per line (a lead of two short sentences reads best this way)."""
        for sentence in sentences:
            assert text_width_in(sentence, size) * 1.04 < w, f"{name}: '{sentence[:40]}...' needs a second line"
        h = lines_height_in(len(sentences), size) + 0.04
        return self.text_lines(s, x, y, w, h, [[N(t, size, colour)] for t in sentences], name=name), h

    def title(self, s, x, y, w, h, text, size, colour=None, lines=None, anchor=MSO_ANCHOR.MIDDLE,
              align=PP_ALIGN.LEFT, spacing=1.0):
        """Fill the layout's title placeholder (kept for screen readers and the outline view). Call it right after
        the decorative background shapes, so the title is the first thing read. lines: hand-set line breaks."""
        colour = colour or P.primary
        for ln in (lines or [text]):
            assert text_width_in(ln, size, True) * 1.05 < w + 1e-6, f"title line too long: {ln!r}"
        assert " ".join(lines or [text]) == text, "title lines must join back to the title"
        ph = s.shapes.title
        ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
        ph.name = "Title"
        ph.fill.background()
        ph.line.fill.background()
        ph.shadow.inherit = False
        tf = ph.text_frame
        self._frame(tf, anchor)
        body = tf._txBody
        for p in body.findall(qn("a:p")):
            body.remove(p)
        para = _Paragraph(etree.SubElement(body, qn("a:p")), tf)
        self._para_format(para, align, spacing)
        for i, ln in enumerate(lines or [text]):
            if i:
                self._break(para, size, True)
            self._runs(para, [(ln, size, True, colour)])
        el = ph._element
        tree = el.getparent()
        tree.remove(el)
        ext = tree.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(el)
        else:
            tree.append(el)
        return ph

    # ---- shapes and lines
    def decorative(self, shape):
        """Mark a shape as decorative, so screen readers skip it."""
        cnv = shape._element.find(".//" + qn("p:cNvPr"))
        ext_lst = cnv.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(cnv, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", DECORATIVE_URI)
        etree.SubElement(ext, "{%s}decorative" % ADEC_NS, nsmap={"adec": ADEC_NS}).set("val", "1")
        return shape

    @staticmethod
    def _pin_line(ln, dash_set):
        """Pin every line property, so the host deck's line style cannot leak in."""
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        if not dash_set and ln.find(qn("a:prstDash")) is None:
            dash = etree.Element(qn("a:prstDash"))
            dash.set("val", "solid")
            ln.find(qn("a:solidFill")).addnext(dash)
        if ln.find(qn("a:round")) is None:
            ln.find(qn("a:prstDash")).addnext(etree.Element(qn("a:round")))

    def shape(self, s, kind, x, y, w, h, fill=None, line=None, line_w=1.0, dash=None, name=None, adj=None,
              decorative=True):
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            sh.name = name
        for i, v in enumerate(adj or []):
            sh.adjustments[i] = v
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(line_w)
            if dash:
                sh.line.dash_style = dash
            self._pin_line(sh._element.spPr.find(qn("a:ln")), bool(dash))
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False                                # an empty effectLst: no theme shadow
        if sh.has_text_frame:
            sh.text_frame.text = ""
        return self.decorative(sh) if decorative else sh

    def box(self, s, x, y, w, h, fill=None, line=None, line_w=1.0, dash=None, r=None, name=None):
        """A rectangle; r = corner radius in inches (the same radius whatever the box size)."""
        if not r:
            return self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, line_w, dash, name)
        return self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, line_w, dash, name,
                          adj=[min(0.5, r / min(w, h))])

    def square(self, s, cx, cy, size, fill=None, name="Square"):
        return self.box(s, cx - size / 2, cy - size / 2, size, size, fill=fill or P.primary, name=name)

    def circle(self, s, cx, cy, d, fill=None, line=None, line_w=1.0, name="Circle"):
        return self.shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill, line, line_w, None, name)

    def line(self, s, x1, y1, x2, y2, colour=None, width_pt=1.0, dash=None, name="Line", arrow_end=False):
        colour = colour or P.primary
        c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.name = name
        c.line.color.rgb = rgb(colour)
        c.line.width = Pt(width_pt)
        if dash:
            c.line.dash_style = dash
        ln = c._element.spPr.find(qn("a:ln"))
        self._pin_line(ln, bool(dash))
        if arrow_end:
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")
        c.shadow.inherit = False                                 # connectors take the theme's shadow otherwise
        return self.decorative(c)

    def arc(self, s, cx, cy, r, start_deg, end_deg, width_pt, colour=None, name="Arc"):
        """A native arc (editable in PowerPoint): angles clockwise from 3 o'clock; 180 to 270 is the upper-left
        quarter. Its frame is the whole circle and may run past the slide edge; only the drawn quarter shows. For a
        decoration that must stay inside a box (a guest slide), use an SVG with svg_picture instead."""
        sh = s.shapes.add_shape(MSO_SHAPE.ARC, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        sh.name = name
        av = sh._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
        for child in list(av):
            av.remove(child)
        for nm, deg in (("adj1", start_deg), ("adj2", end_deg)):
            gd = etree.SubElement(av, qn("a:gd"))
            gd.set("name", nm)
            gd.set("fmla", "val %d" % int(round(deg * 60000)))
        sh.fill.background()
        sh.line.color.rgb = rgb(colour or P.primary)
        sh.line.width = Pt(width_pt)
        self._pin_line(sh._element.spPr.find(qn("a:ln")), False)
        sh.shadow.inherit = False
        return self.decorative(sh)

    def divider(self, s, x1, y, x2, name="Divider"):
        """A dotted hairline across a column."""
        return self.line(s, x1, y, x2, y, P.hairline, 1.0, dash=SQUARE_DOT, name=name)

    def vdivider(self, s, x, y1, y2, name="Divider"):
        """A dotted hairline down a gutter."""
        return self.line(s, x, y1, x, y2, P.hairline, 1.0, dash=SQUARE_DOT, name=name)

    # ---- pictures
    def _svg_part(self, s, svg_bytes):
        key = hashlib.sha1(svg_bytes).hexdigest()
        part = self._svg_parts.get(key)
        if part is None:
            self._svg_count += 1
            part = Part(PackURI(f"/ppt/media/graphic{self._svg_count}.svg"), "image/svg+xml", s.part.package, svg_bytes)
            self._svg_parts[key] = part
        return s.part.relate_to(part, RT.IMAGE)

    def _attach_svg(self, pic, rid):
        blip = pic._element.find(".//" + qn("a:blip"))
        ext = etree.SubElement(etree.SubElement(blip, qn("a:extLst")), qn("a:ext"))
        ext.set("uri", SVG_EXT_URI)
        etree.SubElement(ext, "{%s}svgBlip" % ASVG_NS, nsmap={"asvg": ASVG_NS}).set(qn("r:embed"), rid)

    def icon(self, s, name, x, y, size, colour=None, descr=None):
        """A line icon, inserted the way PowerPoint's Insert > Icons does it (SVG plus a PNG fallback), so it stays
        sharp and can be recoloured in PowerPoint. Icons: assets/icons (Phosphor, MIT licence)."""
        import fitz
        colour = colour or P.primary
        svg = (ICONS / f"{name}.svg").read_text(encoding="utf-8").replace('fill="currentColor"', f'fill="#{colour}"')
        svg_bytes = svg.encode("utf-8")
        page = fitz.open(stream=svg_bytes, filetype="svg")[0]
        px = max(256, int(size * 300))
        png = page.get_pixmap(matrix=fitz.Matrix(px / page.rect.width, px / page.rect.width), alpha=True).tobytes("png")
        pic = s.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), Inches(size), Inches(size))
        pic.name = f"Icon {name}"
        pic._element.find(".//" + qn("p:cNvPr")).set("descr", descr or name.replace("-", " ") + " icon")
        self._attach_svg(pic, self._svg_part(s, svg_bytes))
        return self.decorative(pic)

    def svg_picture(self, s, svg_text, x, y, w, h, name, descr="decorative graphic"):
        """A drawing as SVG with a PNG fallback; anything outside the SVG's viewBox is clipped, so nothing sits on
        the pasteboard off the slide."""
        import fitz
        svg_bytes = svg_text.encode("utf-8")
        page = fitz.open(stream=svg_bytes, filetype="svg")[0]
        px = max(600, int(w * 300))
        png = page.get_pixmap(matrix=fitz.Matrix(px / page.rect.width, px / page.rect.width), alpha=True).tobytes("png")
        pic = s.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), Inches(w), Inches(h))
        pic.name = name
        pic._element.find(".//" + qn("p:cNvPr")).set("descr", descr)
        self._attach_svg(pic, self._svg_part(s, svg_bytes))
        return self.decorative(pic)

    def logo(self, s, path, x, y, w, h, align="left", valign="top", descr="logo"):
        """An official logo, unaltered: cropped to its artwork (the file is untouched), scaled to fit the box,
        never stretched or recoloured. Returns the shape and its displayed box (x, y, w, h)."""
        (dx, dy, dw, dh), crop = self.logo_box(path, x, y, w, h, align, valign)
        pic = s.shapes.add_picture(str(path), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        pic.name = "Logo"
        pic._element.find(".//" + qn("p:cNvPr")).set("descr", descr)
        return pic, (dx, dy, dw, dh)

    @staticmethod
    def logo_box(path, x, y, w, h, align="left", valign="top"):
        """Where logo() would put a logo, and its crop, without adding it (to lay other things out against it)."""
        im = Image.open(path).convert("RGBA")
        bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
        flat = Image.alpha_composite(bg, im).convert("L")
        mask = flat.point(lambda v: 255 if v < 235 else 0)
        alpha = im.getchannel("A").point(lambda v: 255 if v > 20 else 0)
        box = Image.composite(mask, Image.new("L", im.size, 0), alpha).getbbox() or (0, 0) + im.size
        pad = 3
        l0, t0 = max(0, box[0] - pad), max(0, box[1] - pad)
        r0, b0 = min(im.width, box[2] + pad), min(im.height, box[3] + pad)
        aw, ah = r0 - l0, b0 - t0
        scale = min(w / aw, h / ah)
        dw, dh = aw * scale, ah * scale
        dx = x + {"left": 0, "center": (w - dw) / 2, "right": w - dw}[align]
        dy = y + {"top": 0, "middle": (h - dh) / 2, "bottom": h - dh}[valign]
        return (dx, dy, dw, dh), (l0 / im.width, t0 / im.height, 1 - r0 / im.width, 1 - b0 / im.height)

    # ---- speaker notes
    def notes(self, s, paragraphs, size=12):
        """paragraphs: [(text, bold), ...]; () for an empty line. Runs are pinned, so the host's notes master
        cannot change the font or language."""
        tf = s.notes_slide.notes_text_frame
        body = tf._txBody
        for p in body.findall(qn("a:p")):
            body.remove(p)
        for item in paragraphs:
            p = _Paragraph(etree.SubElement(body, qn("a:p")), tf)
            p.space_after = Pt(6)
            if item:
                text, bold = item
                self._runs(p, [(text, size, bold, INK)])

    # ---- save
    def save(self, path, author="", title="", subject=""):
        cp = self.prs.core_properties
        now = datetime.now(timezone.utc).replace(microsecond=0)
        cp.author = cp.last_modified_by = author
        cp.title, cp.subject = title, subject
        cp.created = cp.modified = now
        cp.revision = 1
        cp.keywords = cp.comments = cp.category = ""
        for sl in self.prs.slides:                         # a title placeholder nobody filled shows "Click to add title"
            ph = sl.shapes.title
            if ph is not None and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)
        for part in self.prs.part.package.iter_parts():    # the template's app.xml says 4:3 and 0 slides
            if str(part.partname) == "/docProps/app.xml":
                x = part.blob.decode("utf-8")
                x = re.sub(r"<PresentationFormat>[^<]*</PresentationFormat>", "<PresentationFormat>Widescreen</PresentationFormat>", x)
                x = re.sub(r"<Slides>\d+</Slides>", f"<Slides>{len(self.prs.slides)}</Slides>", x)
                x = re.sub(r"<Notes>\d+</Notes>", f"<Notes>{sum(1 for s in self.prs.slides if s.has_notes_slide)}</Notes>", x)
                part._blob = x.encode("utf-8")
        self.prs.save(str(path))
        return path


def finish(path, slide1_png, author=""):
    """After any PowerPoint save of the file: put back the PNG fallbacks of SVG pictures (PowerPoint's save drops
    them), set the thumbnail from a render of slide 1 (the template's is a blank 4:3 one) and the author fields."""
    import fitz
    prs = Presentation(str(path))
    for sl in prs.slides:
        for blip in sl._element.iter(qn("a:blip")):
            svg = blip.find(".//{%s}svgBlip" % ASVG_NS)
            if svg is None or blip.get(qn("r:embed")):
                continue
            svg_part = sl.part.related_part(svg.get(qn("r:embed")))
            pic = blip.getparent().getparent()
            shown_in = int(pic.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext")).get("cx")) / 914400
            page = fitz.open(stream=svg_part.blob, filetype="svg")[0]
            px = max(256, int(shown_in * 300))
            png = page.get_pixmap(matrix=fitz.Matrix(px / page.rect.width, px / page.rect.width), alpha=True).tobytes("png")
            _, rid = sl.part.get_or_add_image_part(io.BytesIO(png))
            blip.set(qn("r:embed"), rid)
    missing = sum(1 for sl in prs.slides for b in sl._element.iter(qn("a:blip")) if not b.get(qn("r:embed")))
    assert missing == 0, f"{missing} pictures still have no fallback"
    thumb = next((p for p in prs.part.package.iter_parts() if str(p.partname) == "/docProps/thumbnail.jpeg"), None)
    if thumb is not None and slide1_png:
        im = Image.open(slide1_png).convert("RGB")
        im.thumbnail((256, 256), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=90)
        thumb._blob = buf.getvalue()
    if author:
        prs.core_properties.author = prs.core_properties.last_modified_by = author
    prs.save(str(path))
